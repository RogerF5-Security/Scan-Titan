from __future__ import annotations

import html
import json
import os
import re
import shutil
import subprocess
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlparse, urlunparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Side, Border
from openpyxl.utils import get_column_letter


BASE_DIR = Path(os.environ.get("SCAN_TITAN_BASE_DIR", Path(__file__).resolve().parents[2])).resolve()
REPORTS_DIR = Path(os.environ.get("SCAN_TITAN_REPORTS_DIR", BASE_DIR / "reports")).resolve()
DEFAULT_TARGETS_FILE = Path(os.environ.get("SCAN_TITAN_TARGETS_FILE", BASE_DIR / "targets" / "targets.txt")).resolve()
ANSI_RE = re.compile(r"\x1b\[[0-?]*[ -/]*[@-~]")


class ReconMatrixManager:
    """Professional Recon_Matrix.xlsx writer using openpyxl only."""

    SHEET = "Recon Matrix"
    WORDLIST_HITS_SHEET = "Wordlist Hits"
    MAX_WORDLIST_HITS_PER_TARGET = 10000
    AUTOSIZE_SAMPLE_ROWS = 250
    METADATA_ROWS = 4
    HEADER_ROW = 6
    DATA_START_ROW = 7
    METADATA_TEXT = "Matriz automatizada de reconocimiento - Scan Titan Community Edition"
    COLUMNS = [
        "Target",
        "IP",
        "Tipo",
        "URL Base",
        "Puertos Expuestos",
        "Servicios",
        "Tecnologias",
        "Enrutamiento/Endpoints descubiertos",
        "Formularios de Login",
        "Librerias JS",
        "Estado de Cookies",
        "Headers Expuestos",
        "WAF/CDN",
        "WAF wafw00f",
        "Perfil WAF",
        "Pruebas WAF",
        "Cortes Adaptativos",
        "Archivos Expuestos",
        "Subdominios",
        "WebSockets",
        "API Specs",
        "GraphQL",
        "Browser Routes",
        "Browser Storage",
        "Browser Screenshots",
        "Sitemap Sin Autenticacion",
        "Directory Listings",
        "SAML Metadata",
        "Residuos Staticos",
        "Pruebas Diferidas Por Politica",
        "Ultima Ejecucion",
        "WhatWeb Tecnologias",
        "WhatWeb Frameworks",
        "WhatWeb Servidores",
        "WhatWeb Categorias",
        "WhatWeb Estado",
        "WhatWeb Evidencia",
    ]
    WORDLIST_HIT_COLUMNS = [
        "Timestamp",
        "Target",
        "IP",
        "Module",
        "Wordlist",
        "Method",
        "Status",
        "URL",
        "Path",
        "Redirect Location",
        "Content-Type",
        "Size Bytes",
        "Time Seconds",
        "Title",
        "Classification",
        "Soft404 Filtered",
        "Sensitive Marker",
        "Evidence Summary",
        "Notes",
    ]
    MERGE_COLUMNS = {
        "Tecnologias",
        "WhatWeb Tecnologias",
        "WhatWeb Frameworks",
        "WhatWeb Servidores",
        "WhatWeb Categorias",
    }

    def __init__(self, path: Path) -> None:
        self.path = path
        self.path.parent.mkdir(parents=True, exist_ok=True)

    def update_target(self, row: dict[str, Any]) -> None:
        self._update_target(row, preserve_existing=False)

    def update_target_partial(self, row: dict[str, Any]) -> None:
        self._update_target(row, preserve_existing=True)

    def update_target_bundle(
        self,
        row: dict[str, Any],
        target: str,
        hits: list[dict[str, Any]],
        timestamp: str,
    ) -> dict[str, int]:
        """Update target and compact Recon hits with one workbook load/save."""
        workbook = self._load_or_create()
        try:
            self._apply_target_row(workbook, row, preserve_existing=False)
            raw_count = len(hits)
            compacted_hits = self._compact_wordlist_hits(hits)
            worksheet = self._ensure_wordlist_hits_sheet(workbook)
            preserved = self._wordlist_rows_except_target(worksheet, target)
            self._replace_wordlist_rows(worksheet, preserved, target, compacted_hits, timestamp)
            self._style_wordlist_hits_sheet(worksheet)
            self._autosize_columns(
                worksheet,
                self.WORDLIST_HIT_COLUMNS,
                max_rows=self.AUTOSIZE_SAMPLE_ROWS,
            )
            self._save_workbook(workbook)
            return {
                "raw": raw_count,
                "kept": len(compacted_hits),
                "discarded": max(0, raw_count - len(compacted_hits)),
            }
        except Exception:
            workbook.close()
            raise

    def update_wordlist_hits(self, target: str, hits: list[dict[str, Any]], timestamp: str) -> None:
        workbook = self._load_or_create()
        try:
            worksheet = self._ensure_wordlist_hits_sheet(workbook)
            preserved = self._wordlist_rows_except_target(worksheet, target)
            compacted_hits = self._compact_wordlist_hits(hits)
            self._replace_wordlist_rows(worksheet, preserved, target, compacted_hits, timestamp)
            self._style_wordlist_hits_sheet(worksheet)
            self._autosize_columns(
                worksheet,
                self.WORDLIST_HIT_COLUMNS,
                max_rows=self.AUTOSIZE_SAMPLE_ROWS,
            )
            self._save_workbook(workbook)
        except Exception:
            workbook.close()
            raise

    def write_dashboard(self, output_path: Path) -> Path:
        """Build a standalone Recon-only dashboard from the historical matrix."""
        workbook = load_workbook(self.path, data_only=True, read_only=True)
        target_rows = self._dashboard_rows(workbook[self.SHEET], self.COLUMNS)
        hit_sheet = workbook[self.WORDLIST_HITS_SHEET]
        all_hit_rows = self._dashboard_rows(hit_sheet, self.WORDLIST_HIT_COLUMNS)
        hit_total = len(all_hit_rows)
        hit_rows = all_hit_rows[:10000]
        workbook.close()

        unique_ips = {str(row.get("IP") or "").strip() for row in target_rows if str(row.get("IP") or "").strip()}
        open_ports = {
            (str(row.get("IP") or row.get("Target") or ""), item.strip())
            for row in target_rows
            for item in str(row.get("Puertos Expuestos") or "").split("|")
            if item.strip()
        }
        subdomains = {
            item.strip().lower()
            for row in target_rows
            for item in str(row.get("Subdominios") or "").split("|")
            if item.strip()
        }

        cards = [
            ("Activos", len(target_rows)),
            ("IPs", len(unique_ips)),
            ("Puertos observados", len(open_ports)),
            ("Subdominios", len(subdomains)),
            ("Rutas por wordlist", hit_total),
        ]
        target_columns = [
            "Target",
            "IP",
            "Tipo",
            "URL Base",
            "Puertos Expuestos",
            "Servicios",
            "Tecnologias",
            "Enrutamiento/Endpoints descubiertos",
            "WAF/CDN",
            "WAF wafw00f",
            "Perfil WAF",
            "Cortes Adaptativos",
            "Subdominios",
            "WhatWeb Servidores",
            "WhatWeb Frameworks",
            "WhatWeb Tecnologias",
            "WhatWeb Estado",
            "Ultima Ejecucion",
        ]
        hit_columns = [
            "Timestamp",
            "Target",
            "IP",
            "Status",
            "URL",
            "Classification",
            "Soft404 Filtered",
            "Sensitive Marker",
        ]

        page = [
            "<!doctype html><html lang='es'><head><meta charset='utf-8'>",
            "<meta name='viewport' content='width=device-width,initial-scale=1'>",
            "<title>Scan Titan - Recon Dashboard</title>",
            "<style>:root{color-scheme:dark}*{box-sizing:border-box}body{margin:0;background:#07111f;color:#e5edf7;font-family:Segoe UI,Arial,sans-serif}"
            ".wrap{max-width:1800px;margin:auto;padding:28px}.eyebrow{color:#38bdf8;font-weight:700;text-transform:uppercase;letter-spacing:.12em;font-size:12px}"
            "h1{margin:7px 0 4px;font-size:30px}.note{color:#94a3b8;margin:0 0 22px}.cards{display:grid;grid-template-columns:repeat(auto-fit,minmax(160px,1fr));gap:12px;margin:20px 0}"
            ".card{background:#0f1d30;border:1px solid #23364f;border-radius:12px;padding:16px}.label{color:#91a4bd;font-size:12px;text-transform:uppercase}.value{font-size:28px;font-weight:750;margin-top:5px}"
            ".panel{background:#0b1727;border:1px solid #23364f;border-radius:12px;margin-top:18px;overflow:hidden}.panel h2{font-size:17px;margin:0;padding:16px 18px;border-bottom:1px solid #23364f}"
            ".toolbar{padding:12px 18px;border-bottom:1px solid #23364f}input{width:min(620px,100%);background:#07111f;color:#e5edf7;border:1px solid #38506c;border-radius:8px;padding:10px 12px}"
            ".tablebox{overflow:auto;max-height:620px}table{border-collapse:collapse;width:100%;font-size:12px}th,td{padding:9px 10px;border-bottom:1px solid #20334b;border-right:1px solid #20334b;vertical-align:top;white-space:pre-wrap;min-width:110px;max-width:380px;overflow-wrap:anywhere}"
            "th{position:sticky;top:0;background:#122238;color:#dbeafe;text-align:left;z-index:1}tr:hover td{background:#102038}.scope{display:inline-block;background:#0c4a6e;color:#bae6fd;border-radius:999px;padding:5px 10px;font-size:12px;font-weight:700}"
            ".cell-list{display:flex;flex-wrap:wrap;gap:6px;align-items:flex-start}.chip{display:inline-flex;align-items:center;max-width:360px;padding:4px 7px;border:1px solid #31516f;background:#10243a;color:#d8ecff;border-radius:7px;line-height:1.25}.chip.status-200{border-color:#22c55e;color:#bbf7d0}.chip.status-301,.chip.status-302{border-color:#38bdf8;color:#bae6fd}.chip.status-403{border-color:#f59e0b;color:#fde68a}.chip.status-500{border-color:#ef4444;color:#fecaca}.muted{color:#8fa7c1}.mono{font-family:Consolas,monospace}</style></head><body><div class='wrap'>",
            "<div class='eyebrow'>Scan Titan / Inteligencia técnica</div>",
            "<h1>Dashboard de Reconocimiento</h1>",
            "<p class='note'><span class='scope'>RECON SOLAMENTE</span> Inventario de superficie; estos registros no se contabilizan como vulnerabilidades.</p>",
            "<div class='cards'>",
        ]
        for label, value in cards:
            page.append(f"<div class='card'><div class='label'>{html.escape(label)}</div><div class='value'>{value}</div></div>")
        page.extend([
            "</div>",
            self._dashboard_table("Activos y superficie consolidada", "targets", target_columns, target_rows),
            self._dashboard_table(
                f"Rutas observadas por wordlist (muestra {len(hit_rows)} de {hit_total})",
                "hits",
                hit_columns,
                hit_rows,
            ),
            "</div><script>function filterTable(id,q){q=q.toLowerCase();document.querySelectorAll('#'+id+' tbody tr').forEach(r=>r.style.display=r.textContent.toLowerCase().includes(q)?'':'none')}</script></body></html>",
        ])
        output_path.parent.mkdir(parents=True, exist_ok=True)
        temp_path = output_path.with_suffix(output_path.suffix + ".tmp")
        temp_path.write_text("\n".join(page), encoding="utf-8")
        temp_path.replace(output_path)
        return output_path

    def _dashboard_rows(
        self,
        worksheet: Any,
        columns: list[str],
        *,
        limit: int | None = None,
    ) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        for row in worksheet.iter_rows(min_row=self.DATA_START_ROW, values_only=True):
            values = list(row[: len(columns)])
            if not any(value not in (None, "") for value in values):
                continue
            rows.append({column: values[index] if index < len(values) else "" for index, column in enumerate(columns)})
            if limit is not None and len(rows) >= limit:
                break
        return rows

    def _dashboard_table(
        self,
        title: str,
        table_id: str,
        columns: list[str],
        rows: list[dict[str, Any]],
        ) -> str:
        sorted_rows = self._sort_dashboard_rows(table_id, rows)
        header = "".join(f"<th>{html.escape(column)}</th>" for column in columns)
        body: list[str] = []
        for row in sorted_rows:
            cells = "".join(
                f"<td>{self._dashboard_cell_html(column, row.get(column), row)}</td>"
                for column in columns
            )
            body.append(f"<tr>{cells}</tr>")
        if not body:
            body.append(f"<tr><td colspan='{len(columns)}'>Sin datos de reconocimiento.</td></tr>")
        return (
            f"<section class='panel'><h2>{html.escape(title)} <small>({len(sorted_rows)})</small></h2>"
            f"<div class='toolbar'><input aria-label='Filtrar {html.escape(title)}' placeholder='Filtrar esta tabla...' "
            f"oninput=\"filterTable('{table_id}',this.value)\"></div>"
            f"<div class='tablebox'><table id='{table_id}'><thead><tr>{header}</tr></thead>"
            f"<tbody>{''.join(body)}</tbody></table></div></section>"
        )

    def _sort_dashboard_rows(self, table_id: str, rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
        if table_id == "hits":
            status_rank = {"500": 0, "403": 1, "401": 2, "200": 3, "301": 4, "302": 5}
            return sorted(
                rows,
                key=lambda row: (
                    status_rank.get(str(row.get("Status") or ""), 99),
                    str(row.get("Target") or "").casefold(),
                    str(row.get("URL") or "").casefold(),
                ),
            )
        return sorted(rows, key=lambda row: str(row.get("Target") or "").casefold())

    def _dashboard_cell_html(self, column: str, value: Any, row: dict[str, Any]) -> str:
        text = str(value or "").strip()
        if not text:
            return "<span class='muted'>-</span>"
        if column == "Status":
            safe_status = html.escape(text)
            css = f"chip status-{re.sub(r'[^0-9]', '', text)[:3]}"
            return f"<span class='{css}'>{safe_status}</span>"
        if column in {
            "Puertos Expuestos",
            "Servicios",
            "Tecnologias",
            "WhatWeb Tecnologias",
            "WhatWeb Frameworks",
            "WhatWeb Servidores",
            "WhatWeb Categorias",
            "Enrutamiento/Endpoints descubiertos",
            "Subdominios",
            "Headers Expuestos",
            "Librerias JS",
        }:
            items = self._split_dashboard_items(text)
            if not items:
                return html.escape(text)
            chips = []
            for item in items[:80]:
                css = "chip mono" if item.startswith(("http://", "https://", "/")) else "chip"
                chips.append(f"<span class='{css}'>{html.escape(item)}</span>")
            if len(items) > 80:
                chips.append(f"<span class='chip muted'>+{len(items) - 80} mas</span>")
            return f"<div class='cell-list'>{''.join(chips)}</div>"
        return html.escape(text)

    def _split_dashboard_items(self, text: str) -> list[str]:
        parts = re.split(r"\s+\|\s+|\n+", text)
        seen: set[str] = set()
        items: list[str] = []
        for part in parts:
            clean = re.sub(r"\s+", " ", part).strip()
            key = clean.casefold()
            if clean and key not in seen:
                seen.add(key)
                items.append(clean)
        return sorted(items, key=lambda item: item.casefold())

    def run_whatweb_recon(
        self,
        targets_file: Path | str = DEFAULT_TARGETS_FILE,
        *,
        binary: str = "whatweb",
        timeout: int = 90,
    ) -> list[dict[str, Any]]:
        results: list[dict[str, Any]] = []
        for target in read_targets_file(Path(targets_file)):
            result = run_whatweb(target, binary=binary, timeout=timeout)
            results.append(result)
            self.update_target_partial(self._technology_result_to_row(result))
        return results

    def run_wappy_recon(
        self,
        targets_file: Path | str = DEFAULT_TARGETS_FILE,
        *,
        binary: str = "whatweb",
        timeout: int = 90,
    ) -> list[dict[str, Any]]:
        return self.run_whatweb_recon(targets_file, binary=binary, timeout=timeout)

    def _update_target(self, row: dict[str, Any], *, preserve_existing: bool) -> None:
        workbook = self._load_or_create()
        try:
            self._apply_target_row(workbook, row, preserve_existing=preserve_existing)
            self._save_workbook(workbook)
        except Exception:
            workbook.close()
            raise

    def _apply_target_row(self, workbook: Workbook, row: dict[str, Any], *, preserve_existing: bool) -> None:
        worksheet = workbook[self.SHEET]
        target = str(row.get("Target") or "").strip()
        if not target:
            return

        existing_row = self._find_target_row(worksheet, target)
        if existing_row is None:
            existing_row = worksheet.max_row + 1
            if existing_row < self.DATA_START_ROW:
                existing_row = self.DATA_START_ROW

        normalized = {column: self._cell_value(row.get(column, "")) for column in self.COLUMNS}
        normalized["Target"] = target
        normalized["Ultima Ejecucion"] = normalized["Ultima Ejecucion"] or datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        )

        self._ensure_headers(worksheet)
        worksheet["A2"] = self.METADATA_TEXT
        worksheet["A4"] = f"Ultima actualizacion del libro: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        for col_idx, column in enumerate(self.COLUMNS, start=1):
            existing_value = worksheet.cell(row=existing_row, column=col_idx).value
            new_value = normalized.get(column, "")
            value = self._resolved_cell_value(
                column=column,
                existing=existing_value,
                incoming=new_value,
                preserve_existing=preserve_existing,
            )
            cell = worksheet.cell(row=existing_row, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = self._thin_border()

        self._style_workbook(worksheet)
        self._autosize(worksheet)

    def _technology_result_to_row(self, result: dict[str, Any]) -> dict[str, Any]:
        technologies = result.get("technologies", [])
        tech_labels = [format_technology(item) for item in technologies]
        frameworks = [
            format_technology(item)
            for item in technologies
            if _is_framework_category(str(item.get("category", "")))
        ]
        servers = [
            format_technology(item)
            for item in technologies
            if _is_server_category(str(item.get("category", "")), str(item.get("name", "")))
        ]
        categories = sorted({str(item.get("category", "")).strip() for item in technologies if item.get("category")})
        status = result.get("status", "UNKNOWN")
        if result.get("error"):
            status = f"{status}: {result.get('error')}"
        return {
            "Target": result.get("display") or result.get("target") or "",
            "Tipo": "URL",
            "URL Base": result.get("url") or result.get("target") or "",
            "Tecnologias": tech_labels,
            "WhatWeb Tecnologias": tech_labels,
            "WhatWeb Frameworks": frameworks,
            "WhatWeb Servidores": servers,
            "WhatWeb Categorias": categories,
            "WhatWeb Estado": status,
            "WhatWeb Evidencia": result.get("evidence", ""),
            "Ultima Ejecucion": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        }

    def _load_or_create(self) -> Workbook:
        if self.path.exists():
            try:
                workbook = load_workbook(self.path)
                if self.SHEET not in workbook.sheetnames:
                    worksheet = workbook.create_sheet(self.SHEET, 0)
                    self._initialize_sheet(worksheet)
                if self.WORDLIST_HITS_SHEET not in workbook.sheetnames:
                    self._initialize_wordlist_hits_sheet(workbook.create_sheet(self.WORDLIST_HITS_SHEET))
                return workbook
            except Exception:
                pass

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = self.SHEET
        self._initialize_sheet(worksheet)
        self._initialize_wordlist_hits_sheet(workbook.create_sheet(self.WORDLIST_HITS_SHEET))
        return workbook

    def _initialize_sheet(self, worksheet: Any) -> None:
        worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(self.COLUMNS))
        worksheet["A1"] = "SCAN TITAN - MATRIZ PROFESIONAL DE RECONOCIMIENTO"
        worksheet["A2"] = self.METADATA_TEXT
        worksheet["A3"] = "Separacion estricta: esta matriz contiene inteligencia/recon, no vulnerabilidades."
        worksheet["A4"] = f"Ultima actualizacion del libro: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        for row in range(1, 5):
            worksheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(self.COLUMNS))
        self._ensure_headers(worksheet)
        self._style_workbook(worksheet)

    def _ensure_wordlist_hits_sheet(self, workbook: Workbook) -> Any:
        if self.WORDLIST_HITS_SHEET not in workbook.sheetnames:
            return self._initialize_wordlist_hits_sheet(workbook.create_sheet(self.WORDLIST_HITS_SHEET))
        worksheet = workbook[self.WORDLIST_HITS_SHEET]
        self._ensure_wordlist_hit_headers(worksheet)
        return worksheet

    def _initialize_wordlist_hits_sheet(self, worksheet: Any) -> Any:
        worksheet.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=len(self.WORDLIST_HIT_COLUMNS),
        )
        worksheet["A1"] = "SCAN TITAN - RUTAS DETECTADAS POR WORDLIST"
        worksheet["A2"] = self.METADATA_TEXT
        worksheet["A3"] = "Rutas probadas por wordlist con HTTP 200, 301, 302, 403 o 500. Filtrar Soft404 para triage manual."
        worksheet["A4"] = f"Ultima actualizacion del libro: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        for row in range(1, 5):
            worksheet.merge_cells(
                start_row=row,
                start_column=1,
                end_row=row,
                end_column=len(self.WORDLIST_HIT_COLUMNS),
            )
        self._ensure_wordlist_hit_headers(worksheet)
        self._style_wordlist_hits_sheet(worksheet)
        return worksheet

    def _ensure_wordlist_hit_headers(self, worksheet: Any) -> None:
        for col_idx, column in enumerate(self.WORDLIST_HIT_COLUMNS, start=1):
            worksheet.cell(row=self.HEADER_ROW, column=col_idx, value=column)

    def _ensure_headers(self, worksheet: Any) -> None:
        for col_idx, column in enumerate(self.COLUMNS, start=1):
            worksheet.cell(row=self.HEADER_ROW, column=col_idx, value=column)

    def _style_workbook(self, worksheet: Any) -> None:
        title_fill = PatternFill("solid", fgColor="0B1F33")
        subtitle_fill = PatternFill("solid", fgColor="1F2937")
        header_fill = PatternFill("solid", fgColor="111827")
        white = "FFFFFF"

        worksheet["A1"].fill = title_fill
        worksheet["A1"].font = Font(color=white, bold=True, size=14)
        worksheet["A1"].alignment = Alignment(horizontal="center", vertical="center")

        for row in range(2, 5):
            cell = worksheet.cell(row=row, column=1)
            cell.fill = subtitle_fill
            cell.font = Font(color=white, bold=row == 2)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_idx in range(1, len(self.COLUMNS) + 1):
            cell = worksheet.cell(row=self.HEADER_ROW, column=col_idx)
            cell.fill = header_fill
            cell.font = Font(color=white, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = self._thin_border()

        worksheet.freeze_panes = f"A{self.DATA_START_ROW}"
        worksheet.auto_filter.ref = f"A{self.HEADER_ROW}:{get_column_letter(len(self.COLUMNS))}{max(worksheet.max_row, self.HEADER_ROW)}"
        worksheet.row_dimensions[1].height = 24
        worksheet.row_dimensions[self.HEADER_ROW].height = 38

    def _style_wordlist_hits_sheet(self, worksheet: Any) -> None:
        title_fill = PatternFill("solid", fgColor="0B1F33")
        subtitle_fill = PatternFill("solid", fgColor="1F2937")
        header_fill = PatternFill("solid", fgColor="111827")
        white = "FFFFFF"

        worksheet["A1"].fill = title_fill
        worksheet["A1"].font = Font(color=white, bold=True, size=14)
        worksheet["A1"].alignment = Alignment(horizontal="center", vertical="center")

        for row in range(2, 5):
            cell = worksheet.cell(row=row, column=1)
            cell.fill = subtitle_fill
            cell.font = Font(color=white, bold=row == 2)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_idx in range(1, len(self.WORDLIST_HIT_COLUMNS) + 1):
            cell = worksheet.cell(row=self.HEADER_ROW, column=col_idx)
            cell.fill = header_fill
            cell.font = Font(color=white, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = self._thin_border()

        worksheet.freeze_panes = f"A{self.DATA_START_ROW}"
        worksheet.auto_filter.ref = (
            f"A{self.HEADER_ROW}:"
            f"{get_column_letter(len(self.WORDLIST_HIT_COLUMNS))}{max(worksheet.max_row, self.HEADER_ROW)}"
        )
        worksheet.row_dimensions[1].height = 24
        worksheet.row_dimensions[self.HEADER_ROW].height = 38

    def _autosize(self, worksheet: Any) -> None:
        self._autosize_columns(worksheet, self.COLUMNS)

    def _autosize_columns(self, worksheet: Any, columns: list[str], *, max_rows: int | None = None) -> None:
        last_row = worksheet.max_row
        if max_rows is not None:
            last_row = min(last_row, self.HEADER_ROW + max(1, int(max_rows)))
        for col_idx, column in enumerate(columns, start=1):
            max_len = len(column)
            for row_idx in range(1, last_row + 1):
                value = worksheet.cell(row=row_idx, column=col_idx).value
                if value is not None:
                    max_len = max(max_len, min(len(str(value)), 80))
            worksheet.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 16), 70)

    def _wordlist_rows_except_target(self, worksheet: Any, target: str) -> list[list[Any]]:
        target_key = str(target or "").strip().lower()
        rows: list[list[Any]] = []
        for values in worksheet.iter_rows(
            min_row=self.DATA_START_ROW,
            max_col=len(self.WORDLIST_HIT_COLUMNS),
            values_only=True,
        ):
            current_target = str(values[1] or "").strip().lower() if len(values) > 1 else ""
            if current_target and current_target != target_key:
                rows.append(list(values))
        return rows

    def _replace_wordlist_rows(
        self,
        worksheet: Any,
        preserved_rows: list[list[Any]],
        target: str,
        hits: list[dict[str, Any]],
        timestamp: str,
    ) -> None:
        data_rows = max(0, worksheet.max_row - self.DATA_START_ROW + 1)
        if data_rows:
            worksheet.delete_rows(self.DATA_START_ROW, data_rows)
        for values in preserved_rows:
            self._append_wordlist_values(worksheet, values)
        for hit in hits:
            self._append_wordlist_hit(worksheet, target, hit, timestamp)

    def _compact_wordlist_hits(self, hits: list[dict[str, Any]]) -> list[dict[str, Any]]:
        unique: dict[tuple[str, str, str], dict[str, Any]] = {}
        for raw in hits:
            if not isinstance(raw, dict) or raw.get("soft404_filtered"):
                continue
            method = str(raw.get("method") or "GET").upper()
            url = str(raw.get("url") or raw.get("path") or "").strip()
            status = str(raw.get("status") or "").strip()
            if not url:
                continue
            key = (method, url.lower(), status)
            candidate = dict(raw)
            existing = unique.get(key)
            if existing is None or self._wordlist_hit_priority(candidate) > self._wordlist_hit_priority(existing):
                unique[key] = candidate
        ordered = sorted(unique.values(), key=self._wordlist_hit_priority, reverse=True)
        return ordered[: self.MAX_WORDLIST_HITS_PER_TARGET]

    def _wordlist_hit_priority(self, hit: dict[str, Any]) -> tuple[int, int, str, str]:
        status = str(hit.get("status") or "")
        status_rank = {"500": 6, "403": 5, "401": 4, "200": 3, "405": 2}.get(status, 1)
        sensitive = 1 if hit.get("sensitive_marker") else 0
        timestamp = str(hit.get("timestamp") or "")
        url = str(hit.get("url") or hit.get("path") or "")
        return sensitive, status_rank, timestamp, url

    def _append_wordlist_values(self, worksheet: Any, values: list[Any]) -> None:
        row_idx = max(self.DATA_START_ROW, worksheet.max_row + 1)
        for col_idx, value in enumerate(values[: len(self.WORDLIST_HIT_COLUMNS)], start=1):
            cell = worksheet.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = self._thin_border()

    def _append_wordlist_hit(
        self,
        worksheet: Any,
        target: str,
        hit: dict[str, Any],
        timestamp: str,
    ) -> None:
        row_values = {
            "Timestamp": hit.get("timestamp") or timestamp,
            "Target": target,
            "IP": hit.get("ip", ""),
            "Module": hit.get("module", ""),
            "Wordlist": hit.get("wordlist", ""),
            "Method": hit.get("method", "GET"),
            "Status": hit.get("status", ""),
            "URL": hit.get("url", ""),
            "Path": hit.get("path", ""),
            "Redirect Location": hit.get("redirect_location", ""),
            "Content-Type": hit.get("content_type", ""),
            "Size Bytes": hit.get("size_bytes", ""),
            "Time Seconds": hit.get("time_seconds", ""),
            "Title": hit.get("title", ""),
            "Classification": hit.get("classification", ""),
            "Soft404 Filtered": "YES" if hit.get("soft404_filtered") else "NO",
            "Sensitive Marker": "YES" if hit.get("sensitive_marker") else "NO",
            "Evidence Summary": hit.get("evidence_summary", ""),
            "Notes": hit.get("notes", ""),
        }
        row_idx = worksheet.max_row + 1
        if row_idx < self.DATA_START_ROW:
            row_idx = self.DATA_START_ROW
        for col_idx, column in enumerate(self.WORDLIST_HIT_COLUMNS, start=1):
            cell = worksheet.cell(row=row_idx, column=col_idx, value=self._cell_value(row_values.get(column, "")))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = self._thin_border()

    def _save_workbook(self, workbook: Workbook) -> None:
        temporary = self.path.with_name(f".{self.path.stem}.tmp{self.path.suffix}")
        try:
            workbook.save(temporary)
        finally:
            workbook.close()
        temporary.replace(self.path)

    def _find_target_row(self, worksheet: Any, target: str) -> int | None:
        for row_idx in range(self.DATA_START_ROW, worksheet.max_row + 1):
            value = worksheet.cell(row=row_idx, column=1).value
            if str(value or "").strip().lower() == target.lower():
                return row_idx
        return None

    def _cell_value(self, value: Any) -> str:
        if isinstance(value, dict):
            return " | ".join(f"{key}: {val}" for key, val in value.items() if str(val).strip())
        if isinstance(value, (list, tuple, set)):
            return " | ".join(str(item) for item in value if str(item).strip())
        return str(value or "")

    def _resolved_cell_value(
        self,
        *,
        column: str,
        existing: Any,
        incoming: str,
        preserve_existing: bool,
    ) -> str:
        existing_text = self._cell_value(existing)
        incoming_text = self._cell_value(incoming)
        if not preserve_existing:
            return incoming_text
        if not incoming_text:
            return existing_text
        if not existing_text:
            return incoming_text
        if column in self.MERGE_COLUMNS:
            return self._merge_pipe_values(existing_text, incoming_text)
        return incoming_text

    def _merge_pipe_values(self, existing: str, incoming: str) -> str:
        values: list[str] = []
        seen: set[str] = set()
        for item in [*existing.split("|"), *incoming.split("|")]:
            clean = item.strip()
            key = clean.lower()
            if clean and key not in seen:
                seen.add(key)
                values.append(clean)
        return " | ".join(values)

    def _thin_border(self) -> Border:
        side = Side(style="thin", color="CBD5E1")
        return Border(left=side, right=side, top=side, bottom=side)


def read_targets_file(path: Path = DEFAULT_TARGETS_FILE) -> list[str]:
    if not path.exists():
        return []
    targets: list[str] = []
    seen: set[str] = set()
    for line in path.read_text(encoding="utf-8", errors="ignore").splitlines():
        target = sanitize_target(line)
        if not target:
            continue
        key = target.lower()
        if key not in seen:
            seen.add(key)
            targets.append(target)
    return targets


def sanitize_target(value: str) -> str:
    clean = re.sub(r"[\x00-\x1f\x7f]", "", str(value or "")).strip().strip("'\"")
    if not clean or clean.startswith("#"):
        return ""
    if not clean.startswith(("http://", "https://")):
        clean = f"https://{clean}"
    parsed = urlparse(clean)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        return ""
    parsed = parsed._replace(fragment="")
    path = parsed.path or "/"
    return urlunparse((parsed.scheme, parsed.netloc, path, "", parsed.query, ""))


def run_whatweb(target: str, *, binary: str = "whatweb", timeout: int = 90) -> dict[str, Any]:
    resolved_binary = shutil.which(binary) or binary
    display = display_name(target)
    if not shutil.which(binary) and not Path(binary).exists():
        return {
            "target": target,
            "url": target,
            "display": display,
            "status": "WHATWEB_NOT_FOUND",
            "technologies": [],
            "error": f"Binary not found in PATH: {binary}",
        }
    supports_json = whatweb_supports_json(resolved_binary)
    output_path = REPORTS_DIR / f"whatweb_recon_{safe_slug(display)}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
    if supports_json:
        command = [
            resolved_binary,
            "--aggression=3",
            "--follow-redirect=always",
            "--max-redirects=10",
            "--open-timeout=15",
            "--read-timeout=30",
            "--max-threads=1",
            "--colour=never",
            "--no-errors",
            f"--log-json={output_path}",
            target,
        ]
    else:
        command = [resolved_binary, target]
    try:
        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=timeout,
            check=False,
        )
    except subprocess.TimeoutExpired:
        return {
            "target": target,
            "url": target,
            "display": display,
            "status": "TIMEOUT",
            "technologies": [],
            "error": f"Timeout after {timeout}s",
        }
    except Exception as exc:
        return {
            "target": target,
            "url": target,
            "display": display,
            "status": "WHATWEB_ERROR",
            "technologies": [],
            "error": str(exc),
        }

    parsed = parse_whatweb_json(output_path) if supports_json else []
    if not parsed:
        parsed = parse_whatweb_output(completed.stdout)
    status = "OK" if parsed else "NO_TECH"
    if completed.returncode not in {0, None} and not parsed:
        status = "WHATWEB_ERROR"
    return {
        "target": target,
        "url": target,
        "display": display,
        "status": status,
        "technologies": parsed,
        "evidence": summarize_technologies(parsed),
        "json_export": str(output_path) if supports_json else "",
        "stdout": completed.stdout,
        "stderr": completed.stderr,
        "returncode": completed.returncode,
        "error": clean_error(completed.stderr) if status == "WHATWEB_ERROR" else "",
    }


def run_wappy(target: str, *, binary: str = "whatweb", timeout: int = 90) -> dict[str, Any]:
    return run_whatweb(target, binary=binary, timeout=timeout)


def whatweb_supports_json(binary: str) -> bool:
    try:
        completed = subprocess.run(
            [binary, "--help"],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=12,
            check=False,
        )
    except Exception:
        return False
    text = f"{completed.stdout}\n{completed.stderr}"
    return completed.returncode in {0, None} and "--log-json" in text and "--aggression" in text


def parse_whatweb_json(path: Path) -> list[dict[str, str]]:
    if not path.exists() or path.stat().st_size <= 2:
        return []
    try:
        parsed = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
    except Exception:
        return []
    records = parsed if isinstance(parsed, list) else [parsed] if isinstance(parsed, dict) else []
    technologies: list[dict[str, str]] = []
    seen: set[tuple[str, str, str]] = set()
    for record in records:
        if not isinstance(record, dict):
            continue
        plugins = record.get("plugins") if isinstance(record.get("plugins"), dict) else {}
        for plugin, data in plugins.items():
            name = re.sub(r"\s+", " ", str(plugin or "")).strip()
            if not name:
                continue
            values = data if isinstance(data, dict) else {}
            versions = flatten_whatweb_values(values.get("version") or values.get("versions"))
            strings = flatten_whatweb_values(values.get("string") or values.get("strings"))
            version = ", ".join(versions[:3])
            evidence = ", ".join(strings[:3])
            category = classify_technology(name)
            key = (category.casefold(), name.casefold(), version.casefold())
            if key in seen:
                continue
            seen.add(key)
            technologies.append({"category": category, "name": name, "version": version, "evidence": evidence})
    return sorted(technologies, key=technology_sort_key)


def parse_whatweb_output(output: str) -> list[dict[str, str]]:
    clean_output = ANSI_RE.sub("", output or "")
    technologies: list[dict[str, str]] = []
    seen: set[tuple[str, str, str]] = set()
    pattern = re.compile(
        r"(?:^|,\s+)(?P<name>[A-Za-z][A-Za-z0-9_.+-]+)(?:\[(?P<value>[^\]]+)\])?",
        re.IGNORECASE,
    )
    for match in pattern.finditer(clean_output):
        name = re.sub(r"\s+", " ", match.group("name") or "").strip()
        value = re.sub(r"\s+", " ", match.group("value") or "").strip()
        if not name or name.casefold() in {"http", "https", "ok", "found", "forbidden"}:
            continue
        if "://" in name or name.lower() in {"nil", "none"}:
            continue
        category = classify_technology(name)
        version_match = re.search(r"\b\d+(?:\.\d+){1,3}[A-Za-z0-9._-]*\b", value)
        version = version_match.group(0) if version_match else ""
        key = (category.lower(), name.lower(), version.lower())
        if key in seen:
            continue
        seen.add(key)
        technologies.append({"category": category, "name": name, "version": version, "evidence": value})
    return sorted(technologies, key=technology_sort_key)


def parse_wappy_output(output: str) -> list[dict[str, str]]:
    return parse_whatweb_output(output)


def flatten_whatweb_values(value: Any) -> list[str]:
    values: list[str] = []
    if isinstance(value, dict):
        iterable: Any = value.values()
    elif isinstance(value, (list, tuple, set)):
        iterable = value
    elif value in (None, ""):
        iterable = []
    else:
        iterable = [value]
    for item in iterable:
        if isinstance(item, (dict, list, tuple, set)):
            values.extend(flatten_whatweb_values(item))
            continue
        text = re.sub(r"\s+", " ", str(item or "")).strip()
        if text and text.casefold() not in {"nil", "none", "null"}:
            values.append(text[:180])
    return list(dict.fromkeys(values))


def classify_technology(name: str) -> str:
    lower = name.casefold()
    if any(token in lower for token in ("apache", "nginx", "iis", "httpserver", "openresty", "tomcat", "jetty")):
        return "Web Server"
    if any(token in lower for token in ("cloudflare", "akamai", "fastly", "imperva", "sucuri", "incapsula", "varnish")):
        return "CDN/WAF/Proxy"
    if any(token in lower for token in ("jquery", "react", "angular", "vue", "bootstrap", "script", "javascript")):
        return "JavaScript/UI"
    if any(token in lower for token in ("wordpress", "drupal", "joomla", "moodle", "liferay", "cms")):
        return "CMS"
    if any(token in lower for token in ("php", "asp", "django", "laravel", "rails", "flask", "express", "framework")):
        return "Framework/Language"
    if any(token in lower for token in ("strict-transport", "x-frame", "content-security", "cookie", "httponly")):
        return "Security/Header"
    if any(token in lower for token in ("title", "html", "doctype", "country", "ip", "email")):
        return "Metadata"
    return "Technology"


def technology_sort_key(item: dict[str, str]) -> tuple[int, str, str]:
    rank = {
        "Web Server": 0,
        "Framework/Language": 1,
        "CMS": 2,
        "JavaScript/UI": 3,
        "CDN/WAF/Proxy": 4,
        "Security/Header": 5,
        "Technology": 6,
        "Metadata": 7,
    }
    category = str(item.get("category") or "")
    return rank.get(category, 99), str(item.get("name") or "").casefold(), str(item.get("version") or "").casefold()


def summarize_technologies(technologies: list[dict[str, str]]) -> str:
    return " | ".join(format_technology(item) for item in sorted(technologies, key=technology_sort_key)[:80])


def safe_slug(value: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("_") or "target"


def display_name(target: str) -> str:
    parsed = urlparse(target)
    return parsed.hostname or target


def format_technology(item: dict[str, str]) -> str:
    name = str(item.get("name", "")).strip()
    category = str(item.get("category", "")).strip()
    version = str(item.get("version", "")).strip()
    evidence = str(item.get("evidence", "")).strip()
    label = name
    if version:
        label = f"{label} {version}"
    if category:
        label = f"{category}: {label}"
    if evidence:
        label = f"{label} ({evidence[:120]})"
    return label


def clean_error(value: str) -> str:
    text = ANSI_RE.sub("", value or "")
    text = re.sub(r"\s+", " ", text).strip()
    return text[:350]


def _is_framework_category(category: str) -> bool:
    lower = category.lower()
    return any(token in lower for token in ["framework", "javascript", "ui", "web framework"])


def _is_server_category(category: str, name: str) -> bool:
    lower_category = category.lower()
    lower_name = name.lower()
    server_names = ["apache", "nginx", "iis", "express", "gunicorn", "tomcat", "cloudflare", "heroku"]
    return any(token in lower_category for token in ["web server", "paas", "cdn", "reverse proxy"]) or any(
        token in lower_name for token in server_names
    )


def main() -> int:
    manager = ReconMatrixManager(REPORTS_DIR / "Recon_Matrix.xlsx")
    results = manager.run_whatweb_recon(DEFAULT_TARGETS_FILE)
    print(f"WhatWeb recon processed targets: {len(results)}")
    for result in results:
        print(f"  {result.get('display')}: {result.get('status')} ({len(result.get('technologies', []))} technologies)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
